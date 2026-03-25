"""PowerPoint文档处理器"""
import asyncio
import traceback
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any

try:
    from pptx import Presentation
    from pptx.presentation import Presentation as PptxPresentation
    from pptx.slide import Slide
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from rag_document_parser.models.schemas import (
    ParseRequest,
    ParseResult,
    DocumentMetadata,
    PageContent,
    DocumentStructure,
    Heading,
    Paragraph,
    Table as TableModel,
    ImageInfo,
    BoundingBox,
    PageMetadata,
    FileType,
)
from rag_document_parser.services.exceptions import DependencyError, ParsingFailedError
from .base_processor import BaseOfficeProcessor


class PowerPointDocumentProcessor(BaseOfficeProcessor):
    """PowerPoint文档处理器"""

    def __init__(self):
        """初始化PowerPoint处理器"""
        super().__init__("powerpoint")
        self.supported_extensions = ["pptx", "ppt"]

    async def parse(
        self,
        file_path: Path,
        request: ParseRequest,
        request_id: str
    ) -> ParseResult:
        """解析PowerPoint文档"""
        self.log_info(
            "开始解析PowerPoint文档",
            extra={
                "request_id": request_id,
                "file_path": str(file_path),
            }
        )

        # 检查依赖可用性
        if not PPTX_AVAILABLE:
            raise DependencyError(
                "python-pptx",
                "PowerPoint文档处理需要python-pptx库，请安装python-pptx>=0.6.23"
            )

        try:
            # 异步执行解析（CPU密集型操作使用线程池）
            loop = asyncio.get_event_loop()
            parse_result = await loop.run_in_executor(
                None,
                self._parse_sync,
                file_path,
                request,
                request_id
            )

            self.log_info(
                "PowerPoint文档解析完成",
                extra={
                    "request_id": request_id,
                    "file_path": str(file_path),
                    "slides_count": parse_result.document_metadata.page_count
                }
            )

            return parse_result

        except Exception as exc:
            self.log_error(
                "PowerPoint文档解析失败",
                extra={
                    "request_id": request_id,
                    "file_path": str(file_path),
                    "error": str(exc),
                    "traceback": traceback.format_exc()
                }
            )
            raise ParsingFailedError(
                f"PowerPoint文档解析失败: {str(exc)}",
                details={"engine": "powerpoint", "format": "pptx"}
            )

    def _parse_sync(
        self,
        file_path: Path,
        request: ParseRequest,
        request_id: str
    ) -> ParseResult:
        """同步解析PowerPoint文档"""
        try:
            # 打开PowerPoint演示文稿
            prs = Presentation(file_path)

            # 提取文档元数据
            metadata = self._extract_metadata(prs, file_path)

            # 提取文档内容（每张幻灯片作为一个页面）
            pages_content = self._extract_content(prs, request, request_id)

            # 提取文档结构
            structure = self._extract_structure(prs, pages_content, request_id)

            # 创建解析结果
            result = ParseResult(
                content=pages_content,
                document_metadata=metadata,
                structure=structure,
            )

            return result

        except Exception as e:
            self.log_error(f"PowerPoint文档解析失败: {e}")
            raise

    def _extract_metadata(self, prs: PptxPresentation, file_path: Path) -> DocumentMetadata:
        """提取PowerPoint文档元数据"""
        # 获取文件信息
        file_size = file_path.stat().st_size if file_path.exists() else None

        # 获取幻灯片数量
        slide_count = len(prs.slides)

        # 尝试获取核心属性
        core_properties = prs.core_properties

        return DocumentMetadata(
            title=core_properties.title or file_path.stem,
            author=core_properties.author,
            subject=core_properties.subject,
            keywords=core_properties.keywords.split(",") if core_properties.keywords else [],
            created_date=core_properties.created,
            modified_date=core_properties.modified,
            page_count=slide_count,
            file_size=file_size,
            file_type=FileType.PPTX
        )

    def _extract_content(
        self,
        prs: PptxPresentation,
        request: ParseRequest,
        request_id: str
    ) -> List[PageContent]:
        """提取PowerPoint文档内容（每张幻灯片作为一个页面）"""
        pages_content = []

        for slide_index, slide in enumerate(prs.slides):
            # 提取幻灯片文本
            text_content = self._extract_slide_text(slide)

            # 提取表格
            tables = []
            if request.extract_tables:
                tables = self._extract_tables_from_slide(slide, slide_index, request_id)

            # 提取图片信息
            images = []
            if request.extract_images:
                images = self._extract_images_from_slide(slide, slide_index, request_id)

            # 页面元数据
            page_metadata = PageMetadata(
                dimensions={"width": prs.slide_width, "height": prs.slide_height},
                rotation=0,
                page_number=slide_index + 1
            )

            page_content = PageContent(
                page=slide_index + 1,
                text=text_content,
                tables=tables,
                images=images,
                metadata=page_metadata
            )
            pages_content.append(page_content)

        return pages_content

    def _extract_slide_text(self, slide: Slide) -> str:
        """提取幻灯片文本"""
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text:
                    text_parts.append(text)

            # 处理表格中的文本
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text.strip())

        return "\n".join(text_parts)

    def _extract_tables_from_slide(
        self,
        slide: Slide,
        slide_index: int,
        request_id: str
    ) -> List[TableModel]:
        """从幻灯片中提取表格"""
        tables = []

        try:
            table_index = 0
            for shape in slide.shapes:
                if shape.has_table:
                    table = shape.table
                    rows = []

                    for row in table.rows:
                        row_data = []
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            row_data.append(cell_text)
                        rows.append(row_data)

                    if rows:
                        table_model = TableModel(
                            table_index=table_index,
                            rows=rows,
                            bbox=None  # 需要计算边界框
                        )
                        tables.append(table_model)
                        table_index += 1

        except Exception as exc:
            self.log_warning(
                "PowerPoint表格提取失败",
                extra={
                    "request_id": request_id,
                    "slide": slide_index + 1,
                    "error": str(exc)
                }
            )

        return tables

    def _extract_images_from_slide(
        self,
        slide: Slide,
        slide_index: int,
        request_id: str
    ) -> List[ImageInfo]:
        """从幻灯片中提取图片信息"""
        images = []

        try:
            image_index = 0
            for shape in slide.shapes:
                if shape.shape_type == 13:  # 图片形状类型
                    # 获取图片位置和尺寸
                    left = shape.left
                    top = shape.top
                    width = shape.width
                    height = shape.height

                    image_info = ImageInfo(
                        image_index=image_index,
                        description=f"幻灯片{slide_index + 1}中的图片",
                        bbox=BoundingBox(
                            x1=left,
                            y1=top,
                            x2=left + width,
                            y2=top + height
                        )
                    )
                    images.append(image_info)
                    image_index += 1

        except Exception as exc:
            self.log_warning(
                "PowerPoint图片提取失败",
                extra={
                    "request_id": request_id,
                    "slide": slide_index + 1,
                    "error": str(exc)
                }
            )

        return images

    def _extract_structure(
        self,
        prs: PptxPresentation,
        pages_content: List[PageContent],
        request_id: str
    ) -> DocumentStructure:
        """提取PowerPoint文档结构"""
        headings = []
        paragraphs = []

        try:
            # 提取幻灯片标题作为标题
            for slide_index, slide in enumerate(prs.slides):
                # 查找标题占位符
                for shape in slide.shapes:
                    if hasattr(shape, "is_title") and shape.is_title:
                        if shape.text and shape.text.strip():
                            headings.append(
                                Heading(
                                    level=1,
                                    text=shape.text.strip(),
                                    page=slide_index + 1,
                                    bbox=None
                                )
                            )
                            break

                # 如果没有找到标题，使用第一段文本作为标题
                if not any(h.page == slide_index + 1 for h in headings):
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            text = shape.text.strip()
                            if text:
                                headings.append(
                                    Heading(
                                        level=1,
                                        text=text[:100],  # 截断长文本
                                        page=slide_index + 1,
                                        bbox=None
                                    )
                                )
                                break

            # 提取幻灯片文本内容作为段落
            for page_content in pages_content:
                if page_content.text:
                    # 将文本拆分为段落（按空行分割）
                    text_lines = page_content.text.split('\n')
                    current_paragraph = []
                    for line in text_lines:
                        if line.strip():
                            current_paragraph.append(line.strip())
                        elif current_paragraph:
                            paragraph_text = ' '.join(current_paragraph)
                            paragraphs.append(
                                Paragraph(
                                    text=paragraph_text,
                                    page=page_content.page,
                                    bbox=None
                                )
                            )
                            current_paragraph = []

                    # 添加最后一个段落
                    if current_paragraph:
                        paragraph_text = ' '.join(current_paragraph)
                        paragraphs.append(
                            Paragraph(
                                text=paragraph_text,
                                page=page_content.page,
                                bbox=None
                            )
                        )

        except Exception as exc:
            self.log_warning(
                "PowerPoint文档结构提取失败",
                extra={
                    "request_id": request_id,
                    "error": str(exc)
                }
            )

        return DocumentStructure(
            headings=headings,
            paragraphs=paragraphs
        )