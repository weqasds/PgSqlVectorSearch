"""Word文档处理器"""
import asyncio
from io import BytesIO
import traceback
from pathlib import Path

try:
    import docx
    from docx.document import Document as DocxDocument

    WORD_AVAILABLE = True
except ImportError:  # pragma: no cover - 依赖缺失时在运行期处理
    WORD_AVAILABLE = False

try:
    from PIL import Image
except ImportError:  # pragma: no cover - 依赖缺失时在运行期处理
    Image = None

from rag_document_parser.models.schemas import (
    DocumentMetadata,
    DocumentStructure,
    FileType,
    Heading,
    ImageInfo,
    PageContent,
    PageMetadata,
    Paragraph,
    ParseRequest,
    ParseResult,
    Table as TableModel,
)
from rag_document_parser.services.exceptions import DependencyError, ParsingFailedError

from .base_processor import BaseOfficeProcessor


class WordDocumentProcessor(BaseOfficeProcessor):
    """Word文档处理器。"""

    def __init__(self) -> None:
        super().__init__("word")
        self.supported_extensions = ["docx"]

    async def parse(
        self,
        file_path: Path,
        request: ParseRequest,
        request_id: str,
    ) -> ParseResult:
        self.log_info(
            "开始解析Word文档",
            extra={"request_id": request_id, "file_path": str(file_path)},
        )

        if not WORD_AVAILABLE:
            raise DependencyError(
                "python-docx",
                "Word文档处理需要python-docx库，请安装python-docx>=1.1.0",
            )

        try:
            loop = asyncio.get_running_loop()
            result = await loop.run_in_executor(None, self._parse_sync, file_path, request, request_id)
            self.log_info(
                "Word文档解析完成",
                extra={
                    "request_id": request_id,
                    "file_path": str(file_path),
                    "pages_count": result.document_metadata.page_count,
                },
            )
            return result
        except Exception as exc:
            self.log_error(
                "Word文档解析失败",
                extra={
                    "request_id": request_id,
                    "file_path": str(file_path),
                    "error": str(exc),
                    "traceback": traceback.format_exc(),
                },
            )
            raise ParsingFailedError(
                f"Word文档解析失败: {str(exc)}",
                details={"engine": "word", "format": "docx"},
            ) from exc

    def _parse_sync(
        self,
        file_path: Path,
        request: ParseRequest,
        request_id: str,
    ) -> ParseResult:
        document = docx.Document(file_path)
        metadata = self._extract_metadata(document, file_path)
        pages_content = self._extract_content(document, request, request_id)
        structure = self._extract_structure(document, pages_content, request_id)
        return ParseResult(
            content=pages_content,
            document_metadata=metadata,
            structure=structure,
        )

    def _extract_metadata(self, doc: DocxDocument, file_path: Path) -> DocumentMetadata:
        core_properties = doc.core_properties
        file_size = file_path.stat().st_size if file_path.exists() else None
        estimated_pages = max(1, len(doc.paragraphs) // 30 or 1)
        keywords = core_properties.keywords.split(",") if core_properties.keywords else []

        return DocumentMetadata(
            title=core_properties.title or file_path.stem,
            author=core_properties.author,
            subject=core_properties.subject,
            keywords=[item.strip() for item in keywords if item.strip()],
            created_date=core_properties.created,
            modified_date=core_properties.modified,
            page_count=estimated_pages,
            file_size=file_size,
            file_type=FileType.DOCX,
        )

    def _extract_content(
        self,
        doc: DocxDocument,
        request: ParseRequest,
        request_id: str,
    ) -> list[PageContent]:
        text_content = self._extract_text_from_doc(doc)
        tables = self._extract_tables_from_doc(doc, request_id) if request.extract_tables else []
        images = self._extract_images_from_doc(doc, request_id) if request.extract_images else []

        return [
            PageContent(
                page=1,
                text=text_content,
                tables=tables,
                images=images,
                metadata=PageMetadata(
                    dimensions={"width": 0.0, "height": 0.0},
                    rotation=0,
                    page_number=1,
                ),
            )
        ]

    def _extract_text_from_doc(self, doc: DocxDocument) -> str:
        return "\n".join(
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text and paragraph.text.strip()
        )

    def _extract_tables_from_doc(
        self,
        doc: DocxDocument,
        request_id: str,
    ) -> list[TableModel]:
        tables: list[TableModel] = []
        try:
            for table_index, table in enumerate(doc.tables):
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in table.rows
                ]
                if rows:
                    tables.append(TableModel(table_index=table_index, rows=rows, bbox=None))
        except Exception as exc:
            self.log_warning(
                "Word表格提取失败",
                extra={"request_id": request_id, "error": str(exc)},
            )
        return tables

    def _extract_images_from_doc(
        self,
        doc: DocxDocument,
        request_id: str,
    ) -> list[ImageInfo]:
        images: list[ImageInfo] = []
        seen_parts: set[str] = set()

        try:
            for relation in doc.part.rels.values():
                rel_type = getattr(relation, "reltype", "")
                if "image" not in rel_type:
                    continue

                target_part = getattr(relation, "target_part", None)
                part_name = str(getattr(target_part, "partname", relation.rId))
                if target_part is None or part_name in seen_parts:
                    continue

                seen_parts.add(part_name)
                description = getattr(target_part, "filename", None) or Path(part_name).name
                blob = getattr(target_part, "blob", None)
                if blob and Image is not None:
                    with Image.open(BytesIO(blob)) as image:
                        description = f"{description} ({image.width}x{image.height})"

                images.append(
                    ImageInfo(
                        image_index=len(images),
                        description=description,
                        bbox=None,
                    )
                )
        except Exception as exc:
            self.log_warning(
                "Word图片提取失败",
                extra={"request_id": request_id, "error": str(exc)},
            )

        return images

    def _extract_structure(
        self,
        doc: DocxDocument,
        pages_content: list[PageContent],
        request_id: str,
    ) -> DocumentStructure:
        headings: list[Heading] = []
        paragraphs: list[Paragraph] = []

        try:
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue

                style_name = paragraph.style.name.lower() if paragraph.style else ""
                if "heading" in style_name:
                    level = next((index for index in range(1, 7) if str(index) in style_name), 1)
                    headings.append(Heading(level=level, text=text, page=1, bbox=None))
                else:
                    paragraphs.append(Paragraph(text=text, page=1, bbox=None))
        except Exception as exc:
            self.log_warning(
                "Word文档结构提取失败",
                extra={"request_id": request_id, "error": str(exc)},
            )

        return DocumentStructure(headings=headings, paragraphs=paragraphs)
